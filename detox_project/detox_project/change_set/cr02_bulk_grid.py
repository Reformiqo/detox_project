"""CR-02 — Bulk download & re-upload in item child tables.

DETOX Production Change-Set FRD (CHANGE SET 01), rows CR-02.1 .. CR-02.7,
logic CL-04..CL-06, validation CVAL-04, register CZ-31 / CZ-32 / CZ-51.

The FRD authored this against the `manufacturing_custom` app; on the Detox
bench the manufacturing customisations live in **detox_project**, so every
object here ships from detox_project (module = "Detox Project").

What this module delivers:

  CR-02.1  `allow_bulk_edit = 1` Property Setter on every item/child Table
           field the Field-Specifications sheet lists (native Frappe grid
           Download / Upload CSV). Code-first upsert — see
           `apply_bulk_edit_property_setters`.

  CR-02.4  Enhanced XLSX: whitelisted `download_grid_xlsx` /
           `upload_grid_xlsx` (openpyxl) that carry EVERY field of the
           child DocType (including the custom fields the native CSV can
           miss — cost_center, project, custom_purchase_order,
           custom_purchase_order_item, operation_name, item_type,
           standard_rate, qty_per_unit …), mapped back by fieldname.

  CR-02.5  Upload asks Replace (clear grid first) or Append.

  CR-02.6  Validation pass BEFORE any row is written: invalid item codes,
           blank mandatory Cost Center / Project, unknown UOM, non-existent
           PO line — collected row-wise; nothing is inserted if any row
           fails (CVAL-04).

  CR-02.7  Download wherever the user can READ; Upload only in Draft and
           only with write permission. Standard permission model — NO
           `ignore_permissions` on the user path. The grid buttons are
           supplied by the Client Scripts installed in
           `install_client_scripts` (Upload button hidden when
           docstatus != 0).

Wiring: `install()` is idempotent and self-heals both the Property Setters
and the Client Scripts on every run — call it from
`detox_project.setup.after_migrate` so the fix survives FC redeploys
(the exact line to add is reported to the manager, per task scope).
"""

import io

import frappe
from frappe import _

# ---------------------------------------------------------------------------
# CR-02.1 — the item/child Table fields that get `allow_bulk_edit = 1`.
# Enumerated straight from the FRD "Field Specifications" sheet, section
# "CHILD TABLES — BULK DOWNLOAD / UPLOAD PROPERTY (CR-02)". The two
# subcontracting rows each name two tables ("Service / Supplied Items" and
# "Items / Supplied Items"), so the 12 FRD rows expand to 14 Table fields.
# ---------------------------------------------------------------------------
BULK_EDIT_TABLES = (
	("Production Plan", "custom_fg_items"),        # Table 1 — Finished Goods
	("Production Plan", "custom_operations"),      # Table 2 — Operations & Materials
	("Stock Entry", "items"),
	("Stock Entry", "additional_costs"),
	("Material Request", "items"),
	("Purchase Order", "items"),
	("Purchase Receipt", "items"),
	("Purchase Invoice", "items"),
	("Delivery Note", "items"),
	("Sales Invoice", "items"),
	("Subcontracting Order", "service_items"),
	("Subcontracting Order", "supplied_items"),
	("Subcontracting Receipt", "items"),
	("Subcontracting Receipt", "supplied_items"),
)

# Fieldtypes that carry no exportable value (layout / display / nested table).
_NO_EXPORT_FIELDTYPES = {
	"Section Break", "Column Break", "Tab Break", "HTML", "Button",
	"Heading", "Fold", "Image", "Signature", "Geolocation", "Barcode",
	"Table", "Table MultiSelect",
}

# Frappe internal columns we never want to import from an uploaded file.
_META_FIELDS = {
	"name", "idx", "docstatus", "parent", "parentfield", "parenttype",
	"owner", "creation", "modified", "modified_by", "doctype",
	"_user_tags", "_comments", "_assign", "_liked_by",
}


# ===========================================================================
# Install / self-heal  (call from after_migrate)
# ===========================================================================
def install():
	"""Idempotent self-heal: Property Setters + Client Scripts."""
	apply_bulk_edit_property_setters()
	install_client_scripts()


def apply_bulk_edit_property_setters():
	"""CR-02.1 / CZ-31 — upsert `allow_bulk_edit = 1` on every listed Table
	field that actually exists on this site. Idempotent."""
	touched = []
	for doctype, fieldname in BULK_EDIT_TABLES:
		if not frappe.db.exists("DocType", doctype):
			continue
		if not frappe.get_meta(doctype).get_field(fieldname):
			# field not present on this site (e.g. subcontracting variant) —
			# skip rather than create a dangling PS.
			continue
		_upsert_check_property(doctype, fieldname, "allow_bulk_edit", "1")
		frappe.clear_cache(doctype=doctype)
		touched.append(f"{doctype}.{fieldname}")
	frappe.db.commit()
	print(
		f"[detox_project] CR-02 allow_bulk_edit upserted on "
		f"{len(touched)} table field(s): {touched}"
	)
	return touched


def _upsert_check_property(doctype, fieldname, prop, value):
	"""Create or update a Check-type Property Setter (module=Detox Project)."""
	ps_name = f"{doctype}-{fieldname}-{prop}"
	if frappe.db.exists("Property Setter", ps_name):
		if frappe.db.get_value("Property Setter", ps_name, "value") != value:
			frappe.db.set_value("Property Setter", ps_name, "value", value)
		return
	frappe.get_doc(
		{
			"doctype": "Property Setter",
			"name": ps_name,
			"doctype_or_field": "DocField",
			"doc_type": doctype,
			"field_name": fieldname,
			"property": prop,
			"property_type": "Check",
			"value": value,
			"module": "Detox Project",
		}
	).insert(ignore_permissions=True)


# ===========================================================================
# CR-02.4 — Enhanced XLSX export / import (whitelisted, permission-enforced)
# ===========================================================================
@frappe.whitelist()
def download_grid_xlsx(doctype, docname, tablefield):
	"""CR-02.4 / CR-02.7 — stream an .xlsx of the CURRENT rows of `tablefield`
	on (doctype, docname), one column per child field (custom fields
	included). Available wherever the user can READ the document."""
	doctype = (doctype or "").strip()
	# CR-02.7 — read gate; no ignore_permissions on the user path.
	if not frappe.has_permission(doctype, ptype="read", doc=docname):
		raise frappe.PermissionError(
			_("Not permitted to read {0} {1}").format(doctype, docname)
		)

	child_dt = _get_child_doctype(doctype, tablefield)
	doc = frappe.get_doc(doctype, docname)
	child_meta = frappe.get_meta(child_dt)
	fields = _exportable_fields(child_meta)
	content = _build_xlsx(fields, doc.get(tablefield) or [])

	frappe.response["filename"] = f"{docname}-{tablefield}.xlsx"
	frappe.response["filecontent"] = content
	frappe.response["type"] = "binary"


@frappe.whitelist()
def upload_grid_xlsx(doctype, docname, tablefield, file_url, mode="Replace"):
	"""CR-02.4 / CR-02.5 / CR-02.6 / CR-02.7 — parse an uploaded .xlsx, run the
	validation pass, then Replace or Append the rows and save the document.

	Returns ``{"errors": [...]}`` (nothing written) when any row fails
	validation, else ``{"errors": [], "rows": n, ...}``."""
	doctype = (doctype or "").strip()
	mode = (mode or "Replace").strip()
	if mode not in ("Replace", "Append"):
		frappe.throw(_("Invalid upload mode: {0}").format(mode))

	# CR-02.7 — write gate; no ignore_permissions on the user path.
	if not frappe.has_permission(doctype, ptype="write", doc=docname):
		raise frappe.PermissionError(
			_("Not permitted to write {0} {1}").format(doctype, docname)
		)

	child_dt = _get_child_doctype(doctype, tablefield)
	doc = frappe.get_doc(doctype, docname)

	# CR-02.7 — Upload only in Draft (docstatus == 0).
	if doc.docstatus != 0:
		frappe.throw(
			_("Rows can only be uploaded while the document is in Draft.")
		)

	content = _read_uploaded_file(file_url, doctype, docname)
	parsed = _parse_grid_xlsx(content)
	if not parsed:
		frappe.throw(_("The uploaded file has no data rows."))

	# CR-02.6 / CVAL-04 — validate the WHOLE file before writing anything.
	errors = _validate_grid_rows(
		doctype, tablefield, parsed, company=doc.get("company")
	)
	if errors:
		return {"errors": errors, "rows": 0}

	before = len(doc.get(tablefield) or [])
	added = _apply_rows(doc, tablefield, parsed, mode, frappe.get_meta(child_dt))
	# Standard save — runs the document's own validate() and respects
	# permissions (write perm already checked above). No ignore_permissions.
	doc.save()

	return {
		"errors": [],
		"rows": added,
		"mode": mode,
		"before": before,
		"total": len(doc.get(tablefield) or []),
	}


# ===========================================================================
# Helpers (pure — unit-testable without a saved document)
# ===========================================================================
def _get_child_doctype(doctype, tablefield):
	"""Resolve + guard: `tablefield` MUST be a Table field on `doctype`.
	Prevents a caller from pointing the export/import at an arbitrary
	attribute."""
	df = frappe.get_meta(doctype).get_field(tablefield)
	if not df or df.fieldtype not in ("Table", "Table MultiSelect"):
		frappe.throw(
			_("{0} is not a child table on {1}").format(tablefield, doctype)
		)
	return df.options


def _exportable_fields(child_meta):
	"""Every persisted, non-layout field of the child DocType."""
	return [
		df for df in child_meta.fields
		if df.fieldtype not in _NO_EXPORT_FIELDTYPES
	]


def _build_xlsx(fields, rows):
	"""Header row = fieldnames (bold, frozen); one row per child row.
	`rows` is a list of child Document objects (or dicts)."""
	import openpyxl
	from openpyxl.styles import Font

	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = "rows"

	header = [df.fieldname for df in fields]
	ws.append(header)
	for cell in ws[1]:
		cell.font = Font(bold=True)
	ws.freeze_panes = "A2"

	for row in rows:
		getter = row.get if hasattr(row, "get") else (lambda k: None)
		ws.append([_xlsx_cell(getter(df.fieldname)) for df in fields])

	bio = io.BytesIO()
	wb.save(bio)
	return bio.getvalue()


def _xlsx_cell(value):
	"""Coerce a Frappe value into something openpyxl can write."""
	if value is None:
		return None
	if isinstance(value, (list, dict)):
		return frappe.as_json(value)
	# str / int / float / Decimal / date / datetime are written natively.
	return value


def _parse_grid_xlsx(content):
	"""Parse .xlsx bytes into a list of {fieldname: value} dicts.
	Row 1 is the fieldname header; blank rows are skipped."""
	import openpyxl

	wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
	ws = wb.active
	rows_iter = ws.iter_rows(values_only=True)

	try:
		header = next(rows_iter)
	except StopIteration:
		return []

	# column index -> fieldname (skip blank header cells)
	cols = {
		i: str(h).strip()
		for i, h in enumerate(header)
		if h is not None and str(h).strip()
	}
	if not cols:
		return []

	parsed = []
	for raw in rows_iter:
		rec = {}
		for i, fieldname in cols.items():
			val = raw[i] if i < len(raw) else None
			if isinstance(val, str):
				val = val.strip()
				if val == "":
					val = None
			rec[fieldname] = val
		if any(v is not None for v in rec.values()):
			parsed.append(rec)
	wb.close()
	return parsed


def _validate_grid_rows(doctype, tablefield, rows, company=None):
	"""CR-02.6 / CVAL-04 — collect row-wise errors. Returns a list of
	human-readable strings, one per offending file row. File row numbers
	are 1-based counting the header (so the first data row is Row 2)."""
	child_dt = _get_child_doctype(doctype, tablefield)
	child_meta = frappe.get_meta(child_dt)
	fmap = {df.fieldname: df for df in child_meta.fields}

	item_df = fmap.get("item_code")
	cc_df = fmap.get("cost_center")
	proj_df = fmap.get("project")
	po_item_df = fmap.get("custom_purchase_order_item")
	# editable UOM link fields (read-only ones are auto-fetched on save)
	uom_fields = [
		df.fieldname
		for df in child_meta.fields
		if df.fieldtype == "Link" and df.options == "UOM" and not df.read_only
	]

	errors = []
	for i, row in enumerate(rows):
		file_row = i + 2  # +1 header, +1 to make it 1-based
		row_errs = []

		# --- invalid / missing item code -----------------------------------
		if item_df:
			val = row.get("item_code")
			if not val:
				if item_df.reqd:
					row_errs.append(_("Item Code is required"))
			elif not frappe.db.exists("Item", val):
				row_errs.append(_("Item Code '{0}' does not exist").format(val))
			elif frappe.db.get_value("Item", val, "disabled"):
				row_errs.append(_("Item '{0}' is disabled").format(val))

		# --- Cost Center (blank mandatory + validity) ----------------------
		if cc_df:
			val = row.get("cost_center")
			if not val:
				if cc_df.reqd:
					row_errs.append(_("Cost Center is required"))
			else:
				msg = _cost_center_error(val, company)
				if msg:
					row_errs.append(msg)

		# --- Project (blank mandatory + existence) -------------------------
		if proj_df:
			val = row.get("project")
			if not val:
				if proj_df.reqd:
					row_errs.append(_("Project is required"))
			elif not frappe.db.exists("Project", val):
				row_errs.append(_("Project '{0}' does not exist").format(val))

		# --- unknown UOM ---------------------------------------------------
		for uf in uom_fields:
			val = row.get(uf)
			if val and not frappe.db.exists("UOM", val):
				row_errs.append(_("UOM '{0}' does not exist").format(val))

		# --- non-existent PO line ------------------------------------------
		if po_item_df:
			val = row.get("custom_purchase_order_item")
			if val and not frappe.db.exists("Purchase Order Item", val):
				row_errs.append(
					_("PO line '{0}' does not exist").format(val)
				)

		if row_errs:
			errors.append(_("Row {0}: {1}").format(file_row, "; ".join(row_errs)))

	return errors


def _cost_center_error(cost_center, company=None):
	"""Return an error string if the Cost Center is invalid, else None.
	Mirrors CVAL-13 (company match, not a group node, not disabled)."""
	cc = frappe.db.get_value(
		"Cost Center", cost_center, ["is_group", "disabled", "company"], as_dict=True
	)
	if not cc:
		return _("Cost Center '{0}' does not exist").format(cost_center)
	if cc.is_group:
		return _("Cost Center '{0}' is a group node").format(cost_center)
	if cc.disabled:
		return _("Cost Center '{0}' is disabled").format(cost_center)
	if company and cc.company != company:
		return _("Cost Center '{0}' does not belong to company {1}").format(
			cost_center, company
		)
	return None


def _apply_rows(doc, tablefield, rows, mode, child_meta):
	"""Map parsed rows onto the document by fieldname. Replace clears the
	grid first; Append keeps the existing rows. Returns rows appended."""
	valid = {
		df.fieldname
		for df in child_meta.fields
		if df.fieldtype not in _NO_EXPORT_FIELDTYPES
	}

	if mode == "Replace":
		doc.set(tablefield, [])

	appended = 0
	for row in rows:
		mapped = {
			fn: val
			for fn, val in row.items()
			if fn in valid and fn not in _META_FIELDS
		}
		if not any(v is not None for v in mapped.values()):
			continue
		doc.append(tablefield, mapped)
		appended += 1
	return appended


def _read_uploaded_file(file_url, doctype, docname):
	"""Bytes of the XLSX the user attached in the dialog.

	Access guard (defence in depth — the target doc's write perm is already
	checked): the File must be owned by the current user (they uploaded it in
	the dialog) or attached to this exact document, so a caller cannot read an
	arbitrary File by pointing `file_url` at it. System Managers are exempt."""
	rows = frappe.get_all(
		"File",
		filters={"file_url": file_url},
		fields=["name", "owner", "attached_to_doctype", "attached_to_name"],
		limit=1,
	)
	if not rows:
		frappe.throw(_("Uploaded file not found."))
	f = rows[0]
	owns = f.owner == frappe.session.user
	attached_here = (
		f.attached_to_doctype == doctype and f.attached_to_name == docname
	)
	if not (owns or attached_here or "System Manager" in frappe.get_roles()):
		raise frappe.PermissionError(_("Not permitted to read the uploaded file."))

	content = frappe.get_doc("File", f.name).get_content()
	if isinstance(content, str):
		content = content.encode("latin-1", "backslashreplace")
	return content


# ===========================================================================
# CR-02.4 / CR-02.5 / CR-02.7 — Client Scripts (code-first, real docs).
# Kept file-backed here rather than doctype_js because `bench build`
# silently fails on FC worker containers (see apps/detox.md).
# ===========================================================================
_METHOD_PREFIX = "detox_project.detox_project.change_set.cr02_bulk_grid"

_CR02_GRID_JS_HELPERS = """
// CR-02 (CZ-32) — Enhanced XLSX Download / Upload buttons on item grids.
// Server pass: {prefix}.download_grid_xlsx / .upload_grid_xlsx
window.detox_cr02_add_grid_buttons = function(frm, tables) {{
	// CR-02.7 — a saved document is required to download/upload rows.
	if (frm.is_new()) return;
	var METHOD_DL = '{prefix}.download_grid_xlsx';
	tables.forEach(function(t) {{
		var gf = frm.fields_dict[t.fieldname];
		if (!gf || !gf.grid) return;
		var grid = gf.grid;
		// CR-02.7 — Download available wherever the doc can be read.
		grid.add_custom_button(__('Download Rows (XLSX)'), function() {{
			var url = '/api/method/' + METHOD_DL
				+ '?doctype=' + encodeURIComponent(frm.doctype)
				+ '&docname=' + encodeURIComponent(frm.docname)
				+ '&tablefield=' + encodeURIComponent(t.fieldname);
			window.open(url);
		}});
		// CR-02.7 — Upload only in Draft (docstatus === 0).
		if (frm.doc.docstatus === 0) {{
			grid.add_custom_button(__('Upload Rows (XLSX)'), function() {{
				detox_cr02_upload_dialog(frm, t.fieldname, t.label);
			}});
		}}
	}});
}};

window.detox_cr02_upload_dialog = function(frm, tablefield, label) {{
	var before = (frm.doc[tablefield] || []).length;
	var d = new frappe.ui.Dialog({{
		title: __('Upload Rows (XLSX) — {{0}}', [label]),
		fields: [
			{{fieldtype: 'Select', fieldname: 'mode', label: __('Mode'),
			 options: 'Replace\\nAppend', 'default': 'Replace', reqd: 1,
			 description: __('Replace clears all {{0}} existing row(s) first. Append adds below them.', [before])}},
			{{fieldtype: 'Attach', fieldname: 'file', label: __('XLSX File'), reqd: 1}}
		],
		primary_action_label: __('Upload'),
		primary_action: function(values) {{
			var run = function() {{
				frappe.call({{
					method: '{prefix}.upload_grid_xlsx',
					args: {{doctype: frm.doctype, docname: frm.docname,
						tablefield: tablefield, file_url: values.file, mode: values.mode}},
					freeze: true, freeze_message: __('Validating and importing rows…'),
					callback: function(r) {{
						var res = (r && r.message) || {{}};
						if (res.errors && res.errors.length) {{
							// CR-02.6 — nothing imported; show the row-wise list.
							frappe.msgprint({{
								title: __('Upload rejected — {{0}} row(s) have errors', [res.errors.length]),
								indicator: 'red',
								message: res.errors.map(function(e) {{
									return frappe.utils.escape_html(e);
								}}).join('<br>')
							}});
							return;
						}}
						d.hide();
						frappe.show_alert({{
							message: __('Imported {{0}} row(s) [{{1}}]. Rows: {{2}} → {{3}}',
								[res.rows, res.mode, res.before, res.total]),
							indicator: 'green'
						}}, 7);
						frm.reload_doc();
					}}
				}});
			}};
			// CR-02.5 — Replace asks for explicit confirmation.
			if (values.mode === 'Replace' && before > 0) {{
				frappe.confirm(
					__('Replace will delete the {{0}} existing row(s) before importing. Continue?', [before]),
					run
				);
			}} else {{
				run();
			}}
		}}
	}});
	d.show();
}};
""".format(prefix=_METHOD_PREFIX)


def _cr02_client_script_body(doctype, tables):
	"""Render the Client Script JS for one parent doctype."""
	tables_js = ", ".join(
		"{{fieldname: '{fn}', label: '{lbl}'}}".format(fn=fn, lbl=lbl)
		for fn, lbl in tables
	)
	return (
		_CR02_GRID_JS_HELPERS
		+ "\nfrappe.ui.form.on('%s', {\n" % doctype
		+ "\trefresh: function(frm) {\n"
		+ "\t\tdetox_cr02_add_grid_buttons(frm, [%s]);\n" % tables_js
		+ "\t}\n});\n"
	)


# Parent doctype -> [(table fieldname, human label)] the enhanced buttons
# attach to (CR-02.4 names Production Plan + Stock Entry explicitly).
_CR02_CLIENT_SCRIPTS = {
	"Production Plan-CR02 Bulk Grid": (
		"Production Plan",
		[("custom_fg_items", "Finished Goods"),
		 ("custom_operations", "Operations & Materials")],
	),
	"Stock Entry-CR02 Bulk Grid": (
		"Stock Entry",
		[("items", "Items"),
		 ("additional_costs", "Additional Costs")],
	),
}


def install_client_scripts():
	"""CR-02.4 — upsert the two enhanced-XLSX Client Scripts as real docs
	(module = Detox Project). Idempotent."""
	for name, (doctype, tables) in _CR02_CLIENT_SCRIPTS.items():
		body = _cr02_client_script_body(doctype, tables)
		if frappe.db.exists("Client Script", name):
			cs = frappe.get_doc("Client Script", name)
			changed = False
			if cs.script != body:
				cs.script = body
				changed = True
			if cs.enabled != 1:
				cs.enabled = 1
				changed = True
			if cs.get("module") != "Detox Project":
				cs.module = "Detox Project"
				changed = True
			if changed:
				cs.save(ignore_permissions=True)
		else:
			frappe.get_doc(
				{
					"doctype": "Client Script",
					"name": name,
					"dt": doctype,
					"view": "Form",
					"enabled": 1,
					"module": "Detox Project",
					"script": body,
				}
			).insert(ignore_permissions=True)
	frappe.db.commit()
	print("[detox_project] CR-02 client scripts upserted: "
		+ ", ".join(_CR02_CLIENT_SCRIPTS.keys()))
