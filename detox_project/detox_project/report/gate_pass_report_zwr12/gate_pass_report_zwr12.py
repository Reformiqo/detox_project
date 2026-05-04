import io
import json

import frappe
from frappe import _


# ---------------------------------------------------------------------------
# Script Report entry point
# ---------------------------------------------------------------------------
def execute(filters=None):
	filters = frappe._dict(filters or {})
	_validate_filters(filters)
	columns = get_columns()
	data = get_data(filters)
	return columns, data


def _validate_filters(filters):
	if not filters.get("company"):
		frappe.throw(_("Company is mandatory"))
	if not filters.get("from_date") or not filters.get("to_date"):
		frappe.throw(_("Date (From) and Date (To) are mandatory"))
	if frappe.utils.getdate(filters.to_date) < frappe.utils.getdate(filters.from_date):
		frappe.throw(_("Date (To) cannot be earlier than Date (From)"))
	if filters.get("exit_from") and filters.get("exit_to"):
		if frappe.utils.getdate(filters.exit_to) < frappe.utils.getdate(filters.exit_from):
			frappe.throw(_("Vehicle Exit Date (To) cannot be earlier than (From)"))


# ---------------------------------------------------------------------------
# 47 columns — order and labels per FRD § 4 (header misspellings preserved)
# ---------------------------------------------------------------------------
def get_columns():
	return [
		{"fieldname": "sno", "label": _("S.No"), "fieldtype": "Int", "width": 60},
		{"fieldname": "date", "label": _("Date"), "fieldtype": "Date", "width": 95},
		{"fieldname": "gate_pass_no", "label": _("Gate Pass No"), "fieldtype": "Link",
			"options": "Gate Pass", "width": 150},
		{"fieldname": "sales_order_no", "label": _("Sales Order No"), "fieldtype": "Link",
			"options": "Sales Order", "width": 140},
		{"fieldname": "subscription_no", "label": _("Subscription No"), "fieldtype": "Link",
			"options": "Subscription", "width": 140},
		{"fieldname": "customer_name", "label": _("Customer Name"), "fieldtype": "Data", "width": 200},
		{"fieldname": "manifest_no", "label": _("Manifest No"), "fieldtype": "Data", "width": 110},
		{"fieldname": "manifest_date", "label": _("Manifest Date"), "fieldtype": "Date", "width": 100},
		{"fieldname": "manifest_qty", "label": _("Manifest Qty"), "fieldtype": "Float",
			"width": 100, "precision": 3},
		{"fieldname": "remarks_chemist", "label": _("Remarks (Chemist)"),
			"fieldtype": "Small Text", "width": 180},
		{"fieldname": "remarks_crm", "label": _("Remarks (CRM)"),
			"fieldtype": "Small Text", "width": 180},
		{"fieldname": "waste_inward_date", "label": _("Waste Inward Date"),
			"fieldtype": "Date", "width": 110},
		{"fieldname": "vehicle_no", "label": _("Vehicle No"), "fieldtype": "Data", "width": 110},
		{"fieldname": "transporter_name", "label": _("Transporter name"), "fieldtype": "Link",
			"options": "Supplier", "width": 160},
		{"fieldname": "packing_type", "label": _("Packing Type"), "fieldtype": "Data", "width": 130},
		{"fieldname": "item_code", "label": _("Item Code"), "fieldtype": "Link",
			"options": "Item", "width": 140},
		{"fieldname": "item_name", "label": _("Item Name"), "fieldtype": "Data", "width": 180},
		{"fieldname": "waste_type", "label": _("Waste Type"), "fieldtype": "Data", "width": 110},
		{"fieldname": "qty", "label": _("Qty"), "fieldtype": "Float", "width": 80, "precision": 3},
		{"fieldname": "vehicle_type", "label": _("Vehicle Type"), "fieldtype": "Data", "width": 110},
		{"fieldname": "customer_gross_weight", "label": _("Customer Gross Weight"),
			"fieldtype": "Float", "width": 130, "precision": 3},
		{"fieldname": "customer_tare_weight", "label": _("Customer Tare Weight"),
			"fieldtype": "Float", "width": 130, "precision": 3},
		{"fieldname": "customer_net_weight", "label": _("Customer Net Weight"),
			"fieldtype": "Float", "width": 130, "precision": 3},
		{"fieldname": "company_gross_weight", "label": _("Company Gross Weight"),
			"fieldtype": "Float", "width": 130, "precision": 3},
		# Header spelling per FRD (lowercase 'c' preserved from legacy template)
		{"fieldname": "company_tare_weight", "label": _("company Tare Weight"),
			"fieldtype": "Float", "width": 130, "precision": 3},
		{"fieldname": "company_net_weight", "label": _("Company Net Weight"),
			"fieldtype": "Float", "width": 130, "precision": 3},
		{"fieldname": "item_description", "label": _("Item Description"),
			"fieldtype": "Small Text", "width": 180},
		{"fieldname": "waste_nature", "label": _("Waste Nature"), "fieldtype": "Data", "width": 110},
		{"fieldname": "document_review", "label": _("Document Review"),
			"fieldtype": "Data", "width": 120},
		{"fieldname": "term_card", "label": _("Term Card"), "fieldtype": "Data", "width": 90},
		{"fieldname": "qi_status", "label": _("QI Status"), "fieldtype": "Data", "width": 100},
		{"fieldname": "quality_review", "label": _("Quality Review"),
			"fieldtype": "Data", "width": 120},
		{"fieldname": "weighment_status", "label": _("Weighment Status"),
			"fieldtype": "Data", "width": 120},
		{"fieldname": "qc_disposal_pathway", "label": _("QC Disposal Pathway"),
			"fieldtype": "Data", "width": 140},
		# Header misspelling preserved per FRD
		{"fieldname": "gatepass_status", "label": _("Gatepass Stauts"),
			"fieldtype": "Data", "width": 130},
		# Header misspelling preserved per FRD
		{"fieldname": "vehicle_exit_date", "label": _("Vehcile Exit Date"),
			"fieldtype": "Date", "width": 120},
		{"fieldname": "sales_document_item", "label": _("Sales Document Item"),
			"fieldtype": "Int", "width": 130},
		{"fieldname": "created_by", "label": _("Created by"), "fieldtype": "Link",
			"options": "User", "width": 160},
		{"fieldname": "lr_no", "label": _("LR No"), "fieldtype": "Data", "width": 100},
		{"fieldname": "waste_category", "label": _("Waste Category"),
			"fieldtype": "Data", "width": 130},
		{"fieldname": "project", "label": _("Project"), "fieldtype": "Link",
			"options": "Project", "width": 160},
		{"fieldname": "pcb_id", "label": _("PCB ID"), "fieldtype": "Data", "width": 110},
		{"fieldname": "uom", "label": _("UOM"), "fieldtype": "Data", "width": 60},
		{"fieldname": "rt_office", "label": _("RT office"), "fieldtype": "Data", "width": 120},
		{"fieldname": "district", "label": _("District"), "fieldtype": "Data", "width": 120},
		{"fieldname": "pincode", "label": _("Pincode"), "fieldtype": "Data", "width": 90},
		{"fieldname": "customer_primary_address", "label": _("Customer Primary Address"),
			"fieldtype": "Small Text", "width": 240},
	]


# ---------------------------------------------------------------------------
# Data assembly
# ---------------------------------------------------------------------------
ROW_CAP = 10000


def get_data(filters):
	rows = _fetch_gate_passes(filters)
	if not rows:
		return []

	gp_names = [r.gate_pass_no for r in rows]
	so_names = list({r.sales_order_no for r in rows if r.sales_order_no})
	sub_names = list({r.subscription_no for r in rows if r.subscription_no})
	cust_names = list({r.customer for r in rows if r.customer})

	qi_by_gp = _fetch_qi(gp_names)
	wi_by_gp = _fetch_waste_inward(gp_names)
	so_count_by_so = _fetch_so_item_count(so_names)
	so_uom_by_so = _fetch_so_first_uom(so_names)
	sub_by_name = _fetch_subscriptions(sub_names)
	cust_by_name = _fetch_customers(cust_names)
	addr_by_cust = _fetch_customer_primary_address(cust_names)
	gpi_by_gp = _fetch_first_gp_item(gp_names)
	item_desc_by_code = {}

	# Second pass: enrich item description from Item if QI didn't have one
	item_codes_needed = set()
	for r in rows:
		qi = qi_by_gp.get(r.gate_pass_no)
		if not qi or not qi.get("description"):
			gpi = gpi_by_gp.get(r.gate_pass_no)
			if gpi and gpi.get("item_code"):
				item_codes_needed.add(gpi["item_code"])
	if item_codes_needed:
		for i in frappe.db.sql(
			"SELECT name, description FROM `tabItem` WHERE name IN %s",
			[list(item_codes_needed)],
			as_dict=True,
		):
			item_desc_by_code[i.name] = i.description

	result = []
	for idx, r in enumerate(rows, start=1):
		qi = qi_by_gp.get(r.gate_pass_no) or {}
		gpi = gpi_by_gp.get(r.gate_pass_no) or {}
		wi_date = wi_by_gp.get(r.gate_pass_no)
		sub = sub_by_name.get(r.subscription_no) or {}
		cust = cust_by_name.get(r.customer) or {}
		addr = addr_by_cust.get(r.customer) or {}

		item_code = qi.get("item_code") or gpi.get("item_code")
		item_name = qi.get("item_name") or gpi.get("item_name")
		item_desc = qi.get("description") or item_desc_by_code.get(item_code or "") or ""

		pcb_id = cust.get("custom_pcb_id") or sub.get("custom_pcb_id") or ""

		result.append(
			{
				"sno": idx,
				"date": r.date,
				"gate_pass_no": r.gate_pass_no,
				"sales_order_no": r.sales_order_no,
				"subscription_no": r.subscription_no,
				"customer_name": r.customer_name,
				"manifest_no": r.manifest_no,
				"manifest_date": r.manifest_date,
				"manifest_qty": r.manifest_qty,
				# ABP2-I204: QI has 4 remark fields with two duplicate labels —
				# pick whichever is populated.
				# Chemist: `remarks` (Data) ↔ `custom_remarks_chemist` (Small Text)
				# CRM:     `custom_remarks_crm_copy` (Data) ↔ `custom_remarks_crm` (Small Text)
				"remarks_chemist": (qi.get("custom_remarks_chemist")
					or qi.get("remarks")
					or qi.get("custom_analysis_summary")
					or ""),
				"remarks_crm": (qi.get("custom_remarks_crm")
					or qi.get("custom_remarks_crm_copy")
					or ""),
				"waste_inward_date": wi_date,
				"vehicle_no": r.vehicle_no,
				"transporter_name": r.transporter,
				"packing_type": r.packing_type or r.container_no,
				"item_code": item_code,
				"item_name": item_name,
				"waste_type": r.waste_type,
				"qty": r.manifest_qty,
				"vehicle_type": r.vehicle_type,
				"customer_gross_weight": r.customer_gross_weight,
				"customer_tare_weight": r.customer_tare_weight,
				"customer_net_weight": r.customer_net_weight,
				"company_gross_weight": r.company_gross_weight,
				"company_tare_weight": r.company_tare_weight,
				"company_net_weight": r.company_net_weight,
				"item_description": item_desc,
				"waste_nature": r.waste_nature,
				"document_review": _derive_document_review(r.custom_document_review, r.workflow_state),
				"term_card": r.term_card,
				"qi_status": r.custom_qi_status or qi.get("status") or "",
				"quality_review": r.quality_review,
				"weighment_status": r.weighment_status,
				"qc_disposal_pathway": r.qc_disposal_pathway,
				"gatepass_status": r.workflow_state or r.status,
				"vehicle_exit_date": r.vehicle_exit_date,
				"sales_document_item": so_count_by_so.get(r.sales_order_no, 0),
				"created_by": r.owner,
				"lr_no": r.lr_no,
				"waste_category": r.waste_category,
				"project": r.project,
				"pcb_id": pcb_id,
				"uom": so_uom_by_so.get(r.sales_order_no, ""),
				"rt_office": sub.get("custom_rt_office") or "",
				"district": addr.get("custom_district") or "",
				"pincode": addr.get("pincode") or "",
				"customer_primary_address": addr.get("full_address") or "",
			}
		)

	return result


# ---------------------------------------------------------------------------
# Individual fetchers
# ---------------------------------------------------------------------------
def _fetch_gate_passes(filters):
	conditions = [
		"gp.company = %(company)s",
		"gp.docstatus < 2",
	]
	values = {"company": filters.company, "from_date": filters.from_date, "to_date": filters.to_date}

	if filters.get("gate_pass"):
		gp_list = filters.gate_pass if isinstance(filters.gate_pass, list) else [filters.gate_pass]
		conditions.append("gp.name IN %(gate_pass)s")
		values["gate_pass"] = tuple(gp_list)
	else:
		conditions.append("gp.date BETWEEN %(from_date)s AND %(to_date)s")

	if filters.get("project"):
		proj = filters.project if isinstance(filters.project, list) else [filters.project]
		conditions.append("gp.project IN %(project)s")
		values["project"] = tuple(proj)

	if filters.get("customer"):
		cust = filters.customer if isinstance(filters.customer, list) else [filters.customer]
		conditions.append("gp.customer IN %(customer)s")
		values["customer"] = tuple(cust)

	if filters.get("exit_from"):
		conditions.append("DATE(gp.vehicle_exit_time) >= %(exit_from)s")
		values["exit_from"] = filters.exit_from

	if filters.get("exit_to"):
		conditions.append("DATE(gp.vehicle_exit_time) <= %(exit_to)s")
		values["exit_to"] = filters.exit_to

	gp_meta = frappe.get_meta("Gate Pass")
	has_doc_review = gp_meta.has_field("custom_document_review")
	has_qi_status = gp_meta.has_field("custom_qi_status")

	if filters.get("document_review") and has_doc_review:
		dr = filters.document_review if isinstance(filters.document_review, list) else [filters.document_review]
		conditions.append("gp.custom_document_review IN %(document_review)s")
		values["document_review"] = tuple(dr)

	where_clause = " AND ".join(conditions)

	doc_review_col = "gp.custom_document_review" if has_doc_review else "NULL"
	qi_status_col = "gp.custom_qi_status" if has_qi_status else "NULL"

	data = frappe.db.sql(
		f"""
		SELECT
			gp.name AS gate_pass_no,
			gp.date,
			gp.sales_order AS sales_order_no,
			gp.subscription AS subscription_no,
			gp.customer,
			gp.customer_name,
			gp.manifest_no,
			gp.manifest_date,
			gp.manifest_qty,
			gp.vehicle_no,
			gp.transporter,
			gp.packing_type,
			gp.container_no,
			gp.waste_type,
			gp.waste_nature,
			gp.waste_category,
			gp.term_card,
			{doc_review_col} AS custom_document_review,
			{qi_status_col} AS custom_qi_status,
			gp.quality_review,
			gp.weighment_status,
			gp.qc_disposal_pathway,
			gp.workflow_state,
			gp.status,
			gp.vehicle_type,
			gp.lr_no,
			gp.project,
			gp.owner,
			gp.customer_gross_weight,
			gp.customer_tare_weight,
			gp.customer_net_weight,
			gp.company_gross_weight,
			gp.company_tare_weight,
			gp.company_net_weight,
			DATE(gp.vehicle_exit_time) AS vehicle_exit_date
		FROM `tabGate Pass` gp
		WHERE {where_clause}
		ORDER BY gp.date DESC, gp.name
		LIMIT {ROW_CAP + 1}
		""",
		values,
		as_dict=True,
	)

	if len(data) > ROW_CAP:
		frappe.msgprint(
			_("Result truncated to {0} rows. Tighten your filters to see more.").format(ROW_CAP),
			indicator="orange",
			alert=True,
		)
		data = data[:ROW_CAP]

	return data



def _derive_document_review(custom_value, workflow_state):
	"""ABP2-I204: when `custom_document_review` is at the default 'Pending'
	or unset, derive a sensible review status from the workflow state. Once
	a Gate Pass passes QC Approval the document review is effectively done,
	so the report should display 'Accepted' instead of leaving 'Pending'
	contradicting the workflow status."""
	ACCEPTED_STATES = {
		"QC Approved", "Weighing Complete", "Pending Weight Approval",
		"Vehicle Exited", "Submitted",
	}
	if workflow_state == "Cancelled":
		return "Cancelled"
	if custom_value and custom_value != "Pending":
		return custom_value
	if workflow_state in ACCEPTED_STATES:
		return "Accepted"
	return custom_value or "Pending"


def _fetch_qi(gp_names):
	"""Return latest submitted/saved QI per Gate Pass, keyed by GP name."""
	if not gp_names:
		return {}
	qi_meta = frappe.get_meta("Quality Inspection")
	extra_cols = []
	# ABP2-I204: include all 4 candidate remark fields + the standard `remarks`
	# field, so the report shows whichever one the user wrote into.
	for f in ("custom_remarks_chemist", "custom_remarks_crm",
			  "custom_remarks_crm_copy", "remarks", "custom_analysis_summary"):
		if qi_meta.has_field(f):
			extra_cols.append(f"qi.{f}")
	extra_sql = (", " + ", ".join(extra_cols)) if extra_cols else ""
	rows = frappe.db.sql(
		f"""
		SELECT qi.reference_name AS gp, qi.name, qi.item_code, qi.item_name,
			qi.description, qi.status{extra_sql}
		FROM `tabQuality Inspection` qi
		WHERE qi.reference_type = 'Gate Pass'
			AND qi.reference_name IN %s
			AND qi.docstatus < 2
		ORDER BY qi.modified DESC
		""",
		[gp_names],
		as_dict=True,
	)
	# Keep the first (= latest modified) entry per GP
	out = {}
	for r in rows:
		out.setdefault(r.gp, r)
	return out


def _fetch_waste_inward(gp_names):
	"""Earliest Waste Inward date for each GP.

	ABP2-I218: Waste Inward's field is `date` (relabelled to "Posting Date"
	via Property Setter); there is no `posting_date` column. The query
	previously failed with OperationalError "Unknown column 'posting_date'".
	"""
	if not gp_names:
		return {}
	rows = frappe.db.sql(
		"""
		SELECT gate_pass, MIN(date) AS posting_date
		FROM `tabWaste Inward`
		WHERE gate_pass IN %s AND docstatus = 1
		GROUP BY gate_pass
		""",
		[gp_names],
		as_dict=True,
	)
	return {r.gate_pass: r.posting_date for r in rows}


def _fetch_so_item_count(so_names):
	if not so_names:
		return {}
	rows = frappe.db.sql(
		"""
		SELECT parent AS so, COUNT(*) AS cnt
		FROM `tabSales Order Item`
		WHERE parent IN %s
		GROUP BY parent
		""",
		[so_names],
		as_dict=True,
	)
	return {r.so: r.cnt for r in rows}


def _fetch_so_first_uom(so_names):
	if not so_names:
		return {}
	rows = frappe.db.sql(
		"""
		SELECT soi.parent AS so, soi.uom
		FROM `tabSales Order Item` soi
		WHERE soi.parent IN %s AND soi.idx = 1
		""",
		[so_names],
		as_dict=True,
	)
	return {r.so: r.uom for r in rows}


def _fetch_subscriptions(names):
	if not names:
		return {}
	meta = frappe.get_meta("Subscription")
	fields = ["name"]
	for f in ("custom_rt_office", "custom_pcb_id"):
		if meta.has_field(f):
			fields.append(f)
	if len(fields) == 1:
		return {}
	rows = frappe.db.sql(
		"SELECT {f} FROM `tabSubscription` WHERE name IN %s".format(f=", ".join(fields)),
		[names],
		as_dict=True,
	)
	return {r.name: r for r in rows}


def _fetch_customers(names):
	if not names:
		return {}
	meta = frappe.get_meta("Customer")
	fields = ["name", "customer_primary_address"]
	if meta.has_field("custom_pcb_id"):
		fields.append("custom_pcb_id")
	rows = frappe.db.sql(
		"SELECT {f} FROM `tabCustomer` WHERE name IN %s".format(f=", ".join(fields)),
		[names],
		as_dict=True,
	)
	return {r.name: r for r in rows}


def _fetch_customer_primary_address(cust_names):
	"""Return {customer_name: {pincode, custom_district, full_address}}."""
	if not cust_names:
		return {}
	meta = frappe.get_meta("Address")
	extra = []
	if meta.has_field("custom_district"):
		extra.append("a.custom_district")
	extra_sql = (", " + ", ".join(extra)) if extra else ""
	rows = frappe.db.sql(
		f"""
		SELECT c.name AS customer, a.address_line1, a.address_line2, a.city,
			a.state, a.country, a.pincode{extra_sql}
		FROM `tabCustomer` c
		LEFT JOIN `tabAddress` a ON a.name = c.customer_primary_address
		WHERE c.name IN %s
		""",
		[cust_names],
		as_dict=True,
	)
	out = {}
	for r in rows:
		parts = [p for p in (r.address_line1, r.address_line2, r.city, r.state, r.country, r.pincode) if p]
		out[r.customer] = {
			"pincode": r.pincode,
			"custom_district": r.get("custom_district"),
			"full_address": ", ".join(parts),
		}
	return out


def _fetch_first_gp_item(gp_names):
	"""Fallback item info from Gate Pass Item (idx=1) if QI not available."""
	if not gp_names:
		return {}
	rows = frappe.db.sql(
		"""
		SELECT parent AS gp, item_code, item_name
		FROM `tabGate Pass Item`
		WHERE parent IN %s AND idx = 1
		""",
		[gp_names],
		as_dict=True,
	)
	return {r.gp: r for r in rows}


# ---------------------------------------------------------------------------
# Template-layout XLSX export
# ---------------------------------------------------------------------------
@frappe.whitelist()
def export_zwr12(filters=None):
	"""Generate the ZWR12 template-layout .xlsx (banner rows 1–11, headers row 12, data from row 13)."""
	try:
		from openpyxl import Workbook
		from openpyxl.styles import Alignment, Font
	except ImportError:
		frappe.throw(_("openpyxl is required for ZWR12 export"))

	if isinstance(filters, str):
		filters = json.loads(filters)
	filters = frappe._dict(filters or {})
	_validate_filters(filters)

	columns = get_columns()
	data = get_data(filters)

	wb = Workbook()
	ws = wb.active
	ws.title = "ZWR12"

	arial10 = Font(name="Arial", size=10)
	arial10_bold = Font(name="Arial", size=10, bold=True)

	# Row 1: Report name
	ws["A1"] = "Report Name: -"
	ws["B1"] = "Gate Pass Report - ZWR12"
	ws.merge_cells("B1:D1")

	# Row 2: banner
	ws["A2"] = "Report Filter By:-"
	ws.merge_cells("A2:D2")
	ws["A2"].font = arial10_bold

	# Rows 3–9: filter values
	filter_layout = [
		(3, "Company", filters.get("company")),
		(4, "Project", _join(filters.get("project"))),
		(5, "Gate Pass No", _join(filters.get("gate_pass"))),
		(6, "Date", _daterange(filters.get("from_date"), filters.get("to_date"))),
		(7, "Vehicle Exit Date", _daterange(filters.get("exit_from"), filters.get("exit_to"))),
		(8, "Customer", _join(filters.get("customer"))),
		(9, "Docs status", _join(filters.get("document_review"))),
	]
	for row_idx, label, value in filter_layout:
		ws.cell(row=row_idx, column=2, value=label).font = arial10_bold
		c = ws.cell(row=row_idx, column=3, value=value or "")
		c.font = arial10
		c.alignment = Alignment(wrap_text=True)
		ws.merge_cells(start_row=row_idx, start_column=3, end_row=row_idx, end_column=4)

	# Row 12: headers
	headers = [c["label"] for c in columns]
	for i, h in enumerate(headers, start=1):
		cell = ws.cell(row=12, column=i, value=h)
		cell.font = arial10_bold

	# Row 13+: data
	keys = [c["fieldname"] for c in columns]
	for r_idx, row in enumerate(data, start=13):
		for c_idx, k in enumerate(keys, start=1):
			v = row.get(k)
			# Format dates as ISO; openpyxl handles datetime.date directly
			cell = ws.cell(row=r_idx, column=c_idx, value=v)
			cell.font = arial10

	# Column widths
	for i, col in enumerate(columns, start=1):
		width = max(10, min(40, int((col.get("width") or 100) / 7)))
		ws.column_dimensions[_col_letter(i)].width = width

	buf = io.BytesIO()
	wb.save(buf)
	buf.seek(0)

	frappe.response["filename"] = f"gate_pass_report_zwr12_{frappe.utils.today()}.xlsx"
	frappe.response["filecontent"] = buf.getvalue()
	frappe.response["type"] = "binary"


def _join(val):
	if not val:
		return ""
	if isinstance(val, list):
		return ", ".join(str(v) for v in val)
	return str(val)


def _daterange(d_from, d_to):
	if not d_from and not d_to:
		return ""
	fmt = frappe.utils.formatdate
	return f"{fmt(d_from) if d_from else ''} to {fmt(d_to) if d_to else ''}"


def _col_letter(n):
	"""1→A, 27→AA, etc."""
	s = ""
	while n:
		n, r = divmod(n - 1, 26)
		s = chr(65 + r) + s
	return s
