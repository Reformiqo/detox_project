"""CR-02 — Bulk download & re-upload in item child tables.

DETOX Production Change-Set FRD (CHANGE SET 01), CR-02.1 .. CR-02.7.

Exercises the actual server half of the user flow (the grid buttons call
these whitelisted methods):

  * CR-02.1  every listed Table field carries `allow_bulk_edit = 1`
  * CR-02.4  download_grid_xlsx round-trips through openpyxl and carries
             the custom fields (cost_center / project / operation_name …)
  * CR-02.5  upload Replace vs Append
  * CR-02.6  row-wise validation collects errors and writes nothing on failure
  * CR-02.7  Upload blocked on a submitted doc; the Table-field guard

The end-to-end upload tests run against a throwaway draft Production Plan
(copied from a real submitted plan); the IntegrationTestCase transaction is
rolled back after every test, so no production data is mutated.
"""

import io

import frappe
from frappe.tests import IntegrationTestCase

from detox_project.detox_project.change_set import cr02_bulk_grid as cr02


class TestCR02BulkGrid(IntegrationTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		# self-heal so the test is standalone even before after_migrate wires it
		cr02.install()
		frappe.clear_cache()
		# Locate a Production Plan that actually has operation rows — used as
		# the read/download source and the seed for the draft-copy upload tests.
		cls.source_pp = cls._find_pp_with_operations()

	@staticmethod
	def _find_pp_with_operations():
		rows = frappe.db.sql(
			"""
			SELECT p.name, COUNT(o.name) AS ops
			FROM `tabProduction Plan` p
			JOIN `tabDetox Production Plan Operation` o ON o.parent = p.name
			GROUP BY p.name
			HAVING ops > 0
			ORDER BY ops DESC
			LIMIT 1
			""",
			as_dict=True,
		)
		if not rows:
			raise RuntimeError(
				"No Production Plan with custom_operations rows on this site — "
				"cannot exercise the CR-02 round-trip."
			)
		return rows[0].name

	# ---- CR-02.1 : Property Setters --------------------------------------
	def test_01_allow_bulk_edit_property_setters(self):
		"""Every listed Table field that exists carries allow_bulk_edit=1
		both as a Property Setter and in the live meta."""
		checked = 0
		for doctype, fieldname in cr02.BULK_EDIT_TABLES:
			if not frappe.db.exists("DocType", doctype):
				continue
			meta = frappe.get_meta(doctype)
			df = meta.get_field(fieldname)
			if not df:
				continue
			ps_name = f"{doctype}-{fieldname}-allow_bulk_edit"
			self.assertTrue(
				frappe.db.exists("Property Setter", ps_name),
				f"Property Setter missing: {ps_name}",
			)
			self.assertEqual(
				frappe.db.get_value("Property Setter", ps_name, "value"), "1",
				f"{ps_name} value != 1",
			)
			self.assertEqual(
				int(df.allow_bulk_edit or 0), 1,
				f"{doctype}.{fieldname} allow_bulk_edit not live in meta",
			)
			checked += 1
		# the FRD names 14 table fields; at least the core manufacturing ones
		# must be present on any Detox site.
		self.assertGreaterEqual(checked, 10, "too few bulk-edit tables verified")

	# ---- CR-02.4 : Client Scripts ----------------------------------------
	def test_02_client_scripts_installed(self):
		for name, dt in (
			("Production Plan-CR02 Bulk Grid", "Production Plan"),
			("Stock Entry-CR02 Bulk Grid", "Stock Entry"),
		):
			self.assertTrue(
				frappe.db.exists("Client Script", name), f"missing {name}"
			)
			cs = frappe.get_doc("Client Script", name)
			self.assertEqual(cs.dt, dt)
			self.assertEqual(cs.enabled, 1)
			self.assertEqual(cs.module, "Detox Project")
			self.assertIn("download_grid_xlsx", cs.script)
			self.assertIn("upload_grid_xlsx", cs.script)
			self.assertIn(cr02._METHOD_PREFIX, cs.script)

	# ---- CR-02.4 : download → parse (round-trip fidelity) ----------------
	def test_03_download_parse_roundtrip(self):
		content = self._download("Production Plan", self.source_pp, "custom_operations")
		self.assertGreater(len(content), 0)
		parsed = cr02._parse_grid_xlsx(content)
		doc = frappe.get_doc("Production Plan", self.source_pp)
		self.assertEqual(len(parsed), len(doc.custom_operations))
		# custom fields survive the round trip
		self.assertEqual(parsed[0]["item_code"], doc.custom_operations[0].item_code)
		self.assertEqual(parsed[0]["cost_center"], doc.custom_operations[0].cost_center)
		self.assertEqual(parsed[0]["project"], doc.custom_operations[0].project)
		self.assertIn("operation_name", parsed[0])
		# real rows validate clean
		errs = cr02._validate_grid_rows(
			"Production Plan", "custom_operations", parsed, company=doc.company
		)
		self.assertEqual(errs, [], f"real rows should validate clean: {errs}")

	# ---- CR-02.5 : mapper Replace vs Append ------------------------------
	def test_04_mapper_replace_and_append(self):
		content = self._download("Production Plan", self.source_pp, "custom_operations")
		parsed = cr02._parse_grid_xlsx(content)
		child_meta = frappe.get_meta("Detox Production Plan Operation")

		d = frappe.new_doc("Production Plan")
		added = cr02._apply_rows(d, "custom_operations", parsed, "Append", child_meta)
		self.assertEqual(added, len(parsed))
		self.assertEqual(len(d.custom_operations), len(parsed))

		# Append again → doubles
		cr02._apply_rows(d, "custom_operations", parsed, "Append", child_meta)
		self.assertEqual(len(d.custom_operations), 2 * len(parsed))

		# Replace → clears then loads only the given rows
		cr02._apply_rows(d, "custom_operations", parsed[:2], "Replace", child_meta)
		self.assertEqual(len(d.custom_operations), 2)
		self.assertEqual(d.custom_operations[0].item_code, parsed[0]["item_code"])

	# ---- CR-02.6 : validation error collection ---------------------------
	def test_05_validation_collects_row_errors(self):
		content = self._download("Production Plan", self.source_pp, "custom_operations")
		parsed = cr02._parse_grid_xlsx(content)
		self.assertGreaterEqual(len(parsed), 2)

		bad = [dict(parsed[0]), dict(parsed[1])]
		bad[0]["item_code"] = "CR02_NON_EXISTENT_ITEM"   # invalid item
		bad[1]["cost_center"] = None                       # blank mandatory CC

		errs = cr02._validate_grid_rows(
			"Production Plan", "custom_operations", bad,
			company=frappe.db.get_value("Production Plan", self.source_pp, "company"),
		)
		self.assertEqual(len(errs), 2, errs)
		self.assertTrue(errs[0].startswith("Row 2:"))
		self.assertIn("CR02_NON_EXISTENT_ITEM", errs[0])
		self.assertTrue(errs[1].startswith("Row 3:"))
		self.assertIn("Cost Center", errs[1])

	# ---- CR-02.5 : end-to-end whitelisted upload Append + Replace --------
	def test_06_end_to_end_upload_append_then_replace(self):
		draft = self._make_draft_pp()
		orig = len(draft.custom_operations)
		self.assertGreater(orig, 0)

		file_url = self._grid_to_file(draft.name, "custom_operations")

		res = cr02.upload_grid_xlsx(
			"Production Plan", draft.name, "custom_operations", file_url, "Append"
		)
		self.assertEqual(res["errors"], [])
		self.assertEqual(res["mode"], "Append")
		self.assertEqual(res["total"], orig * 2)
		self.assertEqual(
			len(frappe.get_doc("Production Plan", draft.name).custom_operations),
			orig * 2,
		)

		res2 = cr02.upload_grid_xlsx(
			"Production Plan", draft.name, "custom_operations", file_url, "Replace"
		)
		self.assertEqual(res2["errors"], [])
		self.assertEqual(res2["total"], orig)
		self.assertEqual(
			len(frappe.get_doc("Production Plan", draft.name).custom_operations),
			orig,
		)

	# ---- CR-02.6 : a bad row writes NOTHING ------------------------------
	def test_07_bad_upload_writes_nothing(self):
		draft = self._make_draft_pp()
		before = len(draft.custom_operations)

		content = self._download("Production Plan", draft.name, "custom_operations")
		import openpyxl

		wb = openpyxl.load_workbook(io.BytesIO(content))
		ws = wb.active
		header = [c.value for c in ws[1]]
		item_col = header.index("item_code") + 1
		ws.cell(row=2, column=item_col).value = "CR02_NON_EXISTENT_ITEM"
		bio = io.BytesIO()
		wb.save(bio)
		bad_url = self._save_file(bio.getvalue(), "cr02-bad.xlsx", draft.name)

		res = cr02.upload_grid_xlsx(
			"Production Plan", draft.name, "custom_operations", bad_url, "Replace"
		)
		self.assertTrue(res["errors"])
		self.assertEqual(res["rows"], 0)
		# grid untouched
		self.assertEqual(
			len(frappe.get_doc("Production Plan", draft.name).custom_operations),
			before,
		)

	# ---- CR-02.7 : Upload blocked on a submitted document ----------------
	def test_08_upload_blocked_when_submitted(self):
		# source_pp is submitted (docstatus=1); any valid file is fine.
		draft = self._make_draft_pp()
		file_url = self._grid_to_file(draft.name, "custom_operations")
		self.assertEqual(
			frappe.db.get_value("Production Plan", self.source_pp, "docstatus"), 1
		)
		with self.assertRaises(frappe.ValidationError):
			cr02.upload_grid_xlsx(
				"Production Plan", self.source_pp, "custom_operations",
				file_url, "Append",
			)

	# ---- CR-02.7 : the Table-field guard (no arbitrary attribute) --------
	def test_09_tablefield_guard_rejects_non_table(self):
		with self.assertRaises(frappe.ValidationError):
			cr02._get_child_doctype("Production Plan", "company")  # not a Table

	# ------------------------------------------------------------------
	# helpers
	# ------------------------------------------------------------------
	def _download(self, doctype, docname, tablefield):
		frappe.local.response = frappe._dict()
		cr02.download_grid_xlsx(doctype, docname, tablefield)
		content = frappe.response["filecontent"]
		frappe.local.response = frappe._dict()
		return content

	def _make_draft_pp(self):
		src = frappe.get_doc("Production Plan", self.source_pp)
		new = frappe.copy_doc(src)
		# frappe.copy_doc keeps docstatus when frappe.in_test (document.py:2115),
		# so a copy of a submitted plan comes back submitted — force it to Draft.
		new.docstatus = 0
		new.insert(ignore_permissions=True)   # test fixture, not the user path
		return new

	def _grid_to_file(self, docname, tablefield):
		content = self._download("Production Plan", docname, tablefield)
		return self._save_file(content, f"{docname}-{tablefield}.xlsx", docname)

	def _save_file(self, content, filename, docname):
		f = frappe.get_doc(
			{
				"doctype": "File",
				"file_name": filename,
				"attached_to_doctype": "Production Plan",
				"attached_to_name": docname,
				"content": content,
				"is_private": 1,
			}
		).insert(ignore_permissions=True)
		return f.file_url
